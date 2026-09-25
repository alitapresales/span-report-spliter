from copy import deepcopy
from pathlib import Path
import re

from docx import Document
import pytest

from app.core.config import Config
from app.core.detection import detect, NS
from app.core.package import Package
from app.core.personalization import personalize, text, replace_text, key_from_cell, label


def test_replacement_across_runs_preserves_formatting():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run('Project SP').bold = True
    paragraph.add_run('AN 1-').italic = True
    paragraph.add_run('156 end')
    replace_text(paragraph._p, re.compile(r'SPAN 1-156'), 'MENARA 012')
    assert paragraph.text == 'Project MENARA 012 end'
    assert paragraph.runs[0].bold and paragraph.runs[1].italic


@pytest.mark.parametrize('profile', ['clearance', 'sag', 'tower_inclination', 'phase_spacing', 'sections', 'tower_height'])
def test_real_preamble_scope_and_totals(profile):
    paths = [p for p in Path('doc-source/report1').glob('*.docx') if Config.for_source(p).profile == profile]
    if not paths:
        pytest.skip('Local sample not available')
    config = Config.for_source(paths[0])
    with Package(paths[0]) as package:
        spans, _ = detect(package.body, config)
        original = text(package.root)
        for span in [spans[0], spans[len(spans)//2], spans[-1]]:
            root = package.segment(span, spans[0].start)
            before = deepcopy(root)
            personalize(root, span, spans[0].start, profile)
            body = root.find('w:body', NS)
            assert len(body) == len(before.find('w:body', NS))
            intro = list(body)[:spans[0].start]
            assert label(span) in ''.join(text(n) for n in intro)
            assert 'SPAN 1-156' not in ''.join(text(n) for n in intro)
            expected = tuple(map(int, re.findall(r'\d+', span.key)))
            filtered = []
            totals = None
            for table in intro:
                rows = table.findall('w:tr', NS)
                if not rows:
                    continue
                headers = [' '.join(text(c).split()).lower() for c in rows[0].findall('w:tc', NS)]
                if headers == ['jenis deteksi', 'kritis', 'bahaya', 'other', 'jumlah']:
                    totals = [[int(text(c)) for c in row.findall('w:tc', NS)[1:]] for row in rows[1:]]
                col = next((i for i,h in enumerate(headers) if h in {'span','nomor dan posisi span','nomor menara'}), None)
                if col is None:
                    continue
                for row in rows[1:]:
                    cells = row.findall('w:tc', NS)
                    value = text(cells[col])
                    if value.strip():
                        assert key_from_cell(value, config.entity_kind == 'tower') == expected
                    filtered.append(cells)
            if profile == 'clearance':
                counts = [sum(text(c[-1]).strip().casefold() == severity for c in filtered) for severity in ['kritis','bahaya']]
                assert totals == [counts + [len(filtered)-sum(counts),len(filtered)]] * 2
            # Detail body remains byte-equivalent; only the introduction is changed.
            from lxml import etree
            assert [etree.tostring(n) for n in list(body)[spans[0].start:]] == [etree.tostring(n) for n in list(before.find('w:body',NS))[spans[0].start:]]
        assert text(package.root) == original
