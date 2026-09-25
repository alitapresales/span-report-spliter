from io import BytesIO
from pathlib import Path
import csv
import json
import os
import struct
import subprocess
import sys
import zlib
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.shared import Inches
from lxml import etree
import pytest

from app.core.config import Config
from app.core.package import DOC, NS, R, W, Package, parse, xml
from app.services.processing import digest, inspect, process


def png(seed: int, side: int = 4, noise: bool = False) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data))
    rows = b''.join(b'\x00' + (os.urandom(side * 3) if noise else bytes([seed % 256, 80, 150]) * side) for _ in range(side))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!2I5B', side, side, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b'')


def make_report(path: Path, count: int = 3, image_side: int = 4) -> None:
    doc = Document()
    doc.add_paragraph('Project cover')
    doc.sections[0].header.paragraphs[0].text = 'Engineering header'
    doc.sections[0].footer.paragraphs[0].text = 'Engineering footer'
    for i in range(1, count + 1):
        doc.add_heading(f'Tower {i:03} - Tower {i+1:03}', 1)
        doc.add_paragraph(f'Analysis for {i}').runs[0].bold = True
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f'Table {i}'
        table.cell(0, 1).text = '12.5 m'
        doc.add_picture(BytesIO(png(i, image_side, image_side > 4)), width=Inches(2))
    doc.save(path)


def rewrite(path: Path, transform) -> None:
    with ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    transform(contents)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, data in contents.items():
            archive.writestr(name, data)


def test_preservation_and_selective_media(tmp_path):
    source = tmp_path / 'input.docx'
    make_report(source)
    original = digest(source)
    output = tmp_path / 'out'
    summary = process(source, output, Config())
    assert summary['generated'] == summary['detected'] == 3
    assert summary['failed'] == 0
    assert digest(source) == original
    for i, path in enumerate(sorted(output.glob('*.docx')), 1):
        doc = Document(path)
        assert [p.text for p in doc.paragraphs] == [f'Tower {i:03} - Tower {i+1:03}', f'Analysis for {i}', '']
        assert doc.paragraphs[0].style.name == 'Heading 1'
        assert doc.paragraphs[1].runs[0].bold
        assert doc.tables[0].cell(0, 0).text == f'Table {i}'
        assert len(doc.inline_shapes) == 1
        assert doc.inline_shapes[0].width == Inches(2)
        assert doc.sections[0].header.paragraphs[0].text == 'Engineering header'
        assert doc.sections[0].footer.paragraphs[0].text == 'Engineering footer'
        with ZipFile(path) as archive, ZipFile(source) as original_zip:
            media = [name for name in archive.namelist() if name.startswith('word/media/')]
            assert len(media) == 1
            assert archive.read(media[0]) == png(i)
            assert archive.read('word/styles.xml') == original_zip.read('word/styles.xml')
    assert len(list(csv.DictReader((output / 'processing_log.csv').open()))) == 3
    assert len(list(csv.DictReader((output / 'validation_report.csv').open()))) == 3


@pytest.mark.parametrize('title,key', [
    ('Span 001', 'Span_001'), ('SPAN 001', 'Span_001'), ('Span No. 001', 'Span_001'),
    ('TOWER 001 - TOWER 002', 'T001-T002'), ('Tower001-Tower002', 'T001-T002'),
    ('T001-T002', 'T001-T002'), ('T.001 – T.002', 'T001-T002'), ('001-002', 'T001-T002'),
])
def test_normalization(tmp_path, title, key):
    path = tmp_path / 'input.docx'
    doc = Document()
    doc.add_paragraph(title)
    doc.save(path)
    assert inspect(path, Config())['spans'] == [key]


def test_rules_toc_preamble_and_closing(tmp_path):
    path = tmp_path / 'input.docx'
    doc = Document()
    doc.add_paragraph('Cover')
    toc = doc.add_paragraph('Span 999')
    props = toc._p.get_or_add_pPr()
    style = etree.SubElement(props, f'{{{W}}}pStyle')
    style.set(f'{{{W}}}val', 'TOC1')
    doc.add_heading('Span 001', 1)
    doc.add_paragraph('Span 002')  # A prose paragraph must not trigger heading-only detection.
    doc.add_heading('Span 003', 1)
    doc.add_heading('Appendix', 1)
    doc.add_paragraph('Excluded appendix')
    doc.save(path)
    config = Config(heading_only=True, include_preamble=True, closing_patterns=['Appendix'])
    assert inspect(path, config)['spans'] == ['Span_001', 'Span_003']
    process(path, tmp_path / 'out', config)
    text = [p.text for p in Document(tmp_path / 'out/Span_003.docx').paragraphs]
    assert text == ['Cover', 'Span 999', 'Span 003']


def test_section_layout_and_inherited_headers(tmp_path):
    source = tmp_path / 'input.docx'
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'Shared header'
    doc.sections[0].left_margin = Inches(0.7)
    doc.add_heading('Span 001', 1)
    doc.add_paragraph('Portrait')
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    doc.add_heading('Span 002', 1)
    doc.add_paragraph('Landscape')
    doc.save(source)
    process(source, tmp_path / 'out', Config())
    one, two = [Document(tmp_path / f'out/Span_00{i}.docx') for i in (1, 2)]
    assert len(one.sections) == len(two.sections) == 1
    assert one.sections[0].orientation == WD_ORIENT.PORTRAIT
    assert two.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert two.sections[0].page_width == Inches(11)
    assert two.sections[0].left_margin == Inches(0.7)
    assert two.sections[0].header.paragraphs[0].text == 'Shared header'


def test_duplicate_no_overwrite_and_continue(tmp_path):
    source = tmp_path / 'input.docx'
    doc = Document()
    for title in ['Span 001', 'Span 001', 'Span 002']:
        doc.add_heading(title, 1)
    doc.save(source)
    summary = process(source, tmp_path / 'out', Config())
    assert summary['failed'] == 2
    assert summary['generated'] == 1
    assert not (tmp_path / 'out/Span_001.docx').exists()


def test_resume_and_tampered_output(tmp_path):
    source, output = tmp_path / 'input.docx', tmp_path / 'out'
    make_report(source)
    process(source, output, Config())
    target = output / 'T001-T002.docx'
    before = target.stat().st_mtime_ns
    summary = process(source, output, Config(), resume=True)
    assert summary['skipped'] == 3
    assert target.stat().st_mtime_ns == before
    target.write_bytes(b'Changed by user')
    summary = process(source, output, Config(), resume=True)
    assert summary['failed'] == 1 and summary['skipped'] == 2
    assert target.read_bytes() == b'Changed by user'
    target.unlink()
    summary = process(source, output, Config(), resume=True)
    assert summary['generated'] == 1 and summary['skipped'] == 2
    with pytest.raises(ValueError, match='Source/config changed'):
        process(source, output, Config(number_width=4), resume=True)


def test_span_failure_isolated(tmp_path):
    source = tmp_path / 'input.docx'
    make_report(source)
    def break_one(contents):
        root = parse(contents[DOC])
        drawing = root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        drawing.set(f'{{{R}}}embed', 'missingRelationship')
        contents[DOC] = xml(root)
    rewrite(source, break_one)
    summary = process(source, tmp_path / 'out', Config())
    assert summary['failed'] == 1 and summary['generated'] == 2
    assert not list((tmp_path / 'out').glob('.span-*'))


def test_source_path_collision(tmp_path):
    source = tmp_path / 'T001-T002.docx'
    make_report(source)
    original = digest(source)
    with pytest.raises(ValueError, match='source document'):
        process(source, tmp_path, Config())
    assert digest(source) == original


def test_no_spans_and_invalid_docx(tmp_path):
    source = tmp_path / 'input.docx'
    Document().save(source)
    with pytest.raises(ValueError, match='No spans'):
        process(source, tmp_path / 'out', Config())
    assert not (tmp_path / 'out/.processing.lock').exists()
    source.write_text('invalid ZIP')
    result = subprocess.run([sys.executable, 'splitter.py', '--input', str(source), '--analyze'], capture_output=True)
    assert result.returncode == 2


def test_config_validation(tmp_path):
    path = tmp_path / 'rules.yaml'
    path.write_text("span_patterns: ['Line (?P<id>\\d+)']\nheading_only: true\n")
    config = Config.load(path)
    assert config.heading_only
    for kwargs in [{'output_template': '../{span_id}.docx'}, {'span_patterns': ['Span (\\d+)']}, {'heading_only': 'false'}]:
        with pytest.raises(ValueError):
            Config(**kwargs)


def test_cli_120_spans(tmp_path):
    source = tmp_path / 'input.docx'
    make_report(source, 120)
    preview = subprocess.run([sys.executable, 'splitter.py', '--input', str(source), '--analyze'], capture_output=True, text=True)
    assert preview.returncode == 0
    assert json.loads(preview.stdout)['detected'] == 120
    run = subprocess.run([sys.executable, 'splitter.py', '--input', str(source), '--output', str(tmp_path / 'out')], capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    assert json.loads(run.stdout)['generated'] == 120
    assert len(list((tmp_path / 'out').glob('*.docx'))) == 120


@pytest.mark.skipif(os.environ.get('SPAN_STRESS_TEST') != '1', reason='Opt-in 85+ MB / 150-span integration test')
def test_large_report(tmp_path):
    source = tmp_path / 'large.docx'
    make_report(source, 150, image_side=450)
    assert source.stat().st_size > 85_000_000
    summary = process(source, tmp_path / 'out', Config())
    assert summary['generated'] == summary['passed'] == 150
    assert summary['source_unchanged']
    for path in (tmp_path / 'out').glob('*.docx'):
        doc = Document(path)
        assert len(doc.inline_shapes) == len(doc.tables) == 1


def test_recursive_chart_media_and_external_hyperlink(tmp_path):
    source = tmp_path / 'input.docx'
    doc = Document()
    doc.add_heading('Span 001', 1)
    doc.add_heading('Span 002', 1)
    doc.save(source)

    def add_chart(contents):
        rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
        chart_ns = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
        root = parse(contents[DOC])
        first = root.find('w:body/w:p', NS)
        run = etree.SubElement(first, f'{{{W}}}r')
        drawing = etree.SubElement(run, f'{{{W}}}drawing')
        chart = etree.SubElement(drawing, f'{{{chart_ns}}}chart')
        chart.set(f'{{{R}}}id', 'chartTest')
        link = etree.SubElement(first, f'{{{W}}}hyperlink')
        link.set(f'{{{R}}}id', 'externalTest')
        contents[DOC] = xml(root)
        rels = parse(contents['word/_rels/document.xml.rels'])
        for rid, kind, target, external in [
            ('chartTest', 'chart', 'charts/chart1.xml', False),
            ('externalTest', 'hyperlink', 'https://example.com/reference', True),
        ]:
            rel = etree.SubElement(rels, f'{{{rel_ns}}}Relationship', Id=rid, Type=f'{R}/{kind}', Target=target)
            if external:
                rel.set('TargetMode', 'External')
        contents['word/_rels/document.xml.rels'] = xml(rels)
        contents['word/charts/chart1.xml'] = f'<c:chartSpace xmlns:c="{chart_ns}"/>'.encode()
        contents['word/charts/_rels/chart1.xml.rels'] = f'<Relationships xmlns="{rel_ns}"><Relationship Id="nestedImage" Type="{R}/image" Target="../media/chart.png"/></Relationships>'.encode()
        contents['word/media/chart.png'] = png(42)
        types = parse(contents['[Content_Types].xml'])
        ct = 'http://schemas.openxmlformats.org/package/2006/content-types'
        etree.SubElement(types, f'{{{ct}}}Override', PartName='/word/charts/chart1.xml', ContentType='application/vnd.openxmlformats-officedocument.drawingml.chart+xml')
        etree.SubElement(types, f'{{{ct}}}Default', Extension='png', ContentType='image/png')
        contents['[Content_Types].xml'] = xml(types)

    rewrite(source, add_chart)
    summary = process(source, tmp_path / 'out', Config())
    assert summary['generated'] == 2
    with ZipFile(tmp_path / 'out/Span_001.docx') as one, ZipFile(tmp_path / 'out/Span_002.docx') as two:
        assert one.read('word/media/chart.png') == png(42)
        assert 'word/charts/chart1.xml' in one.namelist()
        assert b'https://example.com/reference' in one.read('word/_rels/document.xml.rels')
        assert 'word/charts/chart1.xml' not in two.namelist()
        assert 'word/media/chart.png' not in two.namelist()
        assert b'https://example.com/reference' not in two.read('word/_rels/document.xml.rels')


def test_table_profile_skips_summary_and_keeps_repeated_details(tmp_path):
    source = tmp_path / 'Clearance Danger Detection Report.docx'
    doc = Document()
    summary = doc.add_table(rows=5, cols=2)
    summary.cell(0, 1).text = 'Span'
    for i, value in enumerate(['0999-1000', '0001-0002', '0001-0002', '0002-0003'], 1):
        summary.cell(i, 1).text = value
    doc.add_paragraph('6. Detail Kegentingan')
    for value in ['0001-0002', '0001-0002', '0002-0003']:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 1).text = 'Span'
        table.cell(1, 1).text = value
        doc.add_picture(BytesIO(png(1)))
    doc.save(source)
    config = Config.for_source(source)
    assert inspect(source, config)['spans'] == ['T001-T002', 'T002-T003']
    process(source, tmp_path / 'out', config)
    first = Document(tmp_path / 'out/T001-T002.docx')
    assert len(first.tables) == 3
    assert len(first.inline_shapes) == 2
    assert len(first.tables[0].rows) == 3
    assert all(t.cell(1, 1).text == '0001-0002' for t in first.tables)


def test_sag_profile_groups_phases_but_not_nonconsecutive_duplicates(tmp_path):
    source = tmp_path / 'PowerLine Sag Analysis.docx'
    doc = Document()
    doc.add_paragraph('Lendutan Kawat Details')
    for num, pair in enumerate(['1-2', '1-2', '2-3', '1-2'], 1):
        doc.add_paragraph(f'Tabel {num}. Lendutan {pair}Fasa Kiri Atas')
        doc.add_paragraph(f'Gambar {num}. Lendutan {pair}Fasa Kiri Atas')
    doc.save(source)
    result = inspect(source, Config.for_source(source))
    assert result['detected'] == 3
    assert result['duplicates'] == ['T001-T002']


def test_tower_profile_has_distinct_identity_and_explicit_override(tmp_path):
    source = tmp_path / 'Tower Nominal Height.docx'
    doc = Document()
    doc.add_paragraph('Daftar Tinggi Nominal')
    doc.add_paragraph('Tabel 1. Tinggi Nominal 1')
    doc.add_paragraph('Gambar 1. Tinggi Nominal 1')
    doc.add_paragraph('Tabel 2. Tinggi Nominal 2')
    doc.save(source)
    config = Config.for_source(source)
    assert config.entity_kind == 'tower'
    assert inspect(source, config)['spans'] == ['Tower_001', 'Tower_002']
    assert Config.for_source(source, Path('config/default.yaml')).profile == 'generic'


@pytest.mark.parametrize('profile,count,kind', [
    ('clearance', 95, 'span'), ('sag', 152, 'span'),
    ('tower_inclination', 156, 'tower'), ('phase_spacing', 138, 'span'),
    ('sections', 152, 'span'), ('tower_height', 157, 'tower'),
])
def test_local_report_samples(profile, count, kind):
    samples = Path('doc-source/report1')
    matches = [p for p in samples.glob('*.docx') if Config.for_source(p).profile == profile]
    if not matches:
        pytest.skip('Private local samples are not distributed with the project')
    config = Config.for_source(matches[0])
    result = inspect(matches[0], config)
    assert result['detected'] == count
    assert result['entity_kind'] == kind
    assert not result['duplicates']
